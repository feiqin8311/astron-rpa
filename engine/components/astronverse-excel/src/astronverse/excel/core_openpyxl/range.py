from copy import copy
from datetime import date, datetime, time
from typing import Any, Optional

from astronverse.excel import (
    ClearType,
    FontNameType,
    FontType,
    HorizontalAlign,
    NumberFormatType,
    ReadRangeType,
    SetType,
    VerticalAlign,
)
from astronverse.excel.utils import column_number_to_letter
from openpyxl.comments import Comment
from openpyxl.formula.translate import Translator
from openpyxl.styles import Alignment, Border, Font, PatternFill
from openpyxl.utils import get_column_letter

# get_range_data(use_text=True): COM Text 的近似。整数/浮点按 number_format 做简单格式化
# （General / G/通用格式 / 0.00 / 0.00% / 科学计数 / @）；日期用 strftime。
# 公式单元格：openpyxl 不计算，无缓存值时返回公式字符串。
# 多单元格始终返回二维 list（不像 COM 对单行返回一维 tuple）。
# 颜色：fill.fgColor.rgb（FFRRGGBB）；theme/indexed → (255,255,255)。
# get_range_size：列宽×7、行高（磅）估算 Left/Top/Width/Height，非真实像素。
# 自动填充：等差数列、公式相对引用平移、否则周期复制；不支持星期/月份/等比/flash fill。
# 自适应列宽：按单元格显示文本最大长度估算，无法测量字体。

_H_ALIGN = {
    HorizontalAlign.DEFAULT.value: "general",
    HorizontalAlign.LEFT.value: "left",
    HorizontalAlign.RIGHT.value: "right",
    HorizontalAlign.CENTER.value: "center",
    HorizontalAlign.PADDING.value: "fill",
    HorizontalAlign.BOTH.value: "justify",
    HorizontalAlign.CROSS.value: "centerContinuous",
    HorizontalAlign.DISTRIBUTED.value: "distributed",
}
_V_ALIGN = {
    VerticalAlign.UP.value: "top",
    VerticalAlign.MIDDLE.value: "center",
    VerticalAlign.DOWN.value: "bottom",
    VerticalAlign.BOTH.value: "justify",
    VerticalAlign.DISTRIBUTED.value: "distributed",
}

_DEFAULT_COL_WIDTH = 8.43
_DEFAULT_ROW_HEIGHT = 15.0
_COL_WIDTH_TO_PT = 7.0
_PASTE_TYPES = {
    "all",
    "value_and_format",
    "format",
    "exclude_frame",
    "col_width_only",
    "formula_only",
    "formula_and_format",
    "paste_value",
}


def _get_book(ws):
    from astronverse.excel.core_openpyxl import book_of

    return book_of(ws)


def _mark_dirty(ws) -> None:
    _get_book(ws).dirty = True


def _is_number(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _rgb_hex(rgb: tuple[int, int, int]) -> str:
    return f"FF{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}"


def _color_to_rgb(color) -> tuple[int, int, int]:
    if color is None:
        return 255, 255, 255
    ctype = getattr(color, "type", None)
    if ctype in ("theme", "indexed"):
        return 255, 255, 255
    rgb = getattr(color, "rgb", None)
    if rgb is None:
        return 255, 255, 255
    s = str(rgb)
    if s.startswith("Values."):
        return 255, 255, 255
    if len(s) == 8:
        s = s[2:]
    if len(s) != 6:
        return 255, 255, 255
    try:
        return int(s[0:2], 16), int(s[2:4], 16), int(s[4:6], 16)
    except ValueError:
        return 255, 255, 255


def _format_number(v: float, nf: str) -> str:
    if nf in ("General", "G/通用格式", "0", "general"):
        if isinstance(v, float) and v.is_integer():
            return str(int(v))
        return str(v)
    if nf == "0.00":
        return f"{v:.2f}"
    if nf == "0.00%":
        return f"{v * 100:.2f}%"
    if "E+" in nf or "e+" in nf.lower():
        return f"{v:.2E}"
    if nf == "@":
        return str(v)
    if "¥" in nf or "#" in nf:
        return f"{v:.2f}"
    return str(v)


def _display_text(cell) -> str:
    v = cell.value
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    nf = cell.number_format or "General"
    if isinstance(v, datetime):
        if any(k in nf.lower() for k in ("h", "s")) and "y" not in nf.lower() and "年" not in nf:
            return v.strftime("%H:%M:%S")
        return v.strftime("%Y/%m/%d")
    if isinstance(v, date) and not isinstance(v, datetime):
        return v.strftime("%Y/%m/%d")
    if isinstance(v, time):
        return v.strftime("%H:%M:%S")
    if _is_number(v):
        return _format_number(float(v), nf)
    return str(v)


def _cell_value(cell):
    return cell.value


def _snapshot_cell(cell) -> dict:
    comment = cell.comment
    return {
        "value": cell.value,
        "font": copy(cell.font),
        "fill": copy(cell.fill),
        "border": copy(cell.border),
        "alignment": copy(cell.alignment),
        "number_format": cell.number_format,
        "comment": comment.text if comment else None,
        "comment_author": comment.author if comment else "AstronRPA",
    }


def _clear_formats(cell) -> None:
    cell.font = Font()
    cell.fill = PatternFill()
    cell.border = Border()
    cell.alignment = Alignment()
    cell.number_format = "General"


def _clear_all(cell) -> None:
    cell.value = None
    _clear_formats(cell)
    cell.comment = None


def _apply_snapshot(cell, snap: dict, *, values: bool, formats: bool, borders: bool, comments: bool) -> None:
    if values:
        cell.value = snap.get("value")
    if formats:
        if snap.get("font") is not None:
            cell.font = copy(snap["font"])
        if snap.get("fill") is not None:
            cell.fill = copy(snap["fill"])
        if snap.get("alignment") is not None:
            cell.alignment = copy(snap["alignment"])
        if snap.get("number_format") is not None:
            cell.number_format = snap["number_format"]
        if borders and snap.get("border") is not None:
            cell.border = copy(snap["border"])
        elif not borders:
            cell.border = Border()
    if comments:
        text = snap.get("comment")
        if text:
            cell.comment = Comment(text, snap.get("comment_author") or "AstronRPA")
        else:
            cell.comment = None


def _col_width(ws, col: int) -> float:
    dim = ws.column_dimensions.get(get_column_letter(col))
    if dim is None or dim.width is None:
        return _DEFAULT_COL_WIDTH
    return float(dim.width)


def _row_height(ws, row: int) -> float:
    dim = ws.row_dimensions.get(row)
    if dim is None or dim.height is None:
        return _DEFAULT_ROW_HEIGHT
    return float(dim.height)


def _system_copy(text: str) -> None:
    try:
        import pyperclip

        pyperclip.copy(text)
    except Exception:
        pass


def _system_paste() -> str:
    try:
        import pyperclip

        return pyperclip.paste() or ""
    except Exception:
        return ""


def _to_tsv(grid: list[list]) -> str:
    lines = []
    for row in grid:
        lines.append("\t".join("" if v is None else str(v) for v in row))
    return "\n".join(lines)


def _parse_tsv(text: str) -> list[list[str]]:
    if not text:
        return []
    rows = []
    for line in text.split("\n"):
        rows.append(line.rstrip("\r").split("\t"))
    if rows and rows[-1] == [""]:
        rows.pop()
    return rows


def _grid_from_range(rng) -> list[list]:
    ws = rng.ws
    return [
        [ws.cell(r, c).value for c in range(rng.min_col, rng.max_col + 1)] for r in range(rng.min_row, rng.max_row + 1)
    ]


def _clip_from_range(rng) -> dict:
    ws = rng.ws
    cells = []
    for r in range(rng.min_row, rng.max_row + 1):
        cells.append([_snapshot_cell(ws.cell(r, c)) for c in range(rng.min_col, rng.max_col + 1)])
    return {
        "rows": rng.max_row - rng.min_row + 1,
        "cols": rng.max_col - rng.min_col + 1,
        "cells": cells,
        "col_widths": [_col_width(ws, c) for c in range(rng.min_col, rng.max_col + 1)],
    }


def _clip_from_tsv(text: str) -> dict:
    table = _parse_tsv(text)
    cells = [[{"value": v, "number_format": "General"} for v in row] for row in table]
    rows = len(cells)
    cols = max((len(r) for r in cells), default=0)
    for row in cells:
        while len(row) < cols:
            row.append({"value": None, "number_format": "General"})
    return {"rows": rows, "cols": cols, "cells": cells, "col_widths": [_DEFAULT_COL_WIDTH] * cols}


def _transpose_clip(clip: dict) -> dict:
    cells = clip["cells"]
    if not cells:
        return clip
    cols = max(len(r) for r in cells)
    rows = len(cells)
    out = []
    for c in range(cols):
        out.append([cells[r][c] if c < len(cells[r]) else {"value": None} for r in range(rows)])
    widths = clip.get("col_widths") or []
    return {"rows": cols, "cols": rows, "cells": out, "col_widths": widths[:rows] if widths else []}


def _is_blank(snap: dict) -> bool:
    v = snap.get("value")
    return v is None or v == ""


def _paste_clip(rng, clip: dict, paste_type: str, skip_blanks: bool) -> None:
    ws = rng.ws
    r0, c0 = rng.min_row, rng.min_col
    want_values = paste_type in (
        "all",
        "value_and_format",
        "exclude_frame",
        "formula_only",
        "formula_and_format",
        "paste_value",
    )
    want_formats = paste_type in (
        "all",
        "value_and_format",
        "format",
        "exclude_frame",
        "formula_and_format",
    )
    want_borders = paste_type in ("all", "format")
    want_comments = paste_type in ("all", "exclude_frame")
    formulas_only = paste_type in ("formula_only", "formula_and_format")

    if paste_type == "col_width_only":
        widths = clip.get("col_widths") or []
        for i, w in enumerate(widths):
            ws.column_dimensions[get_column_letter(c0 + i)].width = w
        return

    for i, row in enumerate(clip.get("cells") or []):
        for j, snap in enumerate(row):
            if skip_blanks and _is_blank(snap):
                continue
            if formulas_only:
                val = snap.get("value")
                if not (isinstance(val, str) and val.startswith("=")):
                    if paste_type == "formula_only":
                        continue
            cell = ws.cell(r0 + i, c0 + j)
            _apply_snapshot(
                cell,
                snap,
                values=want_values,
                formats=want_formats,
                borders=want_borders,
                comments=want_comments,
            )


def _shift_formula(formula, from_coord: str, to_coord: str):
    if not (isinstance(formula, str) and formula.startswith("=")):
        return formula
    try:
        return Translator(formula, origin=from_coord).translate_formula(to_coord)
    except Exception:
        return formula


def _delete_shift(rng, rows: int, cols: int) -> None:
    ws = rng.ws
    if rows:
        last = ws.max_row
        if rng.max_row < last:
            src = f"{get_column_letter(rng.min_col)}{rng.max_row + 1}:{get_column_letter(rng.max_col)}{last}"
            ws.move_range(src, rows=-rows, cols=0, translate=True)
            vac0 = last - rows + 1
            for r in range(max(vac0, rng.min_row), last + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    _clear_all(ws.cell(r, c))
        else:
            for cell in rng.cells():
                _clear_all(cell)
    elif cols:
        last = ws.max_column
        if rng.max_col < last:
            src = f"{get_column_letter(rng.max_col + 1)}{rng.min_row}:{get_column_letter(last)}{rng.max_row}"
            ws.move_range(src, rows=0, cols=-cols, translate=True)
            vac0 = last - cols + 1
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(max(vac0, rng.min_col), last + 1):
                    _clear_all(ws.cell(r, c))
        else:
            for cell in rng.cells():
                _clear_all(cell)


def _autofit_column(ws, col: int) -> None:
    max_len = 0
    for r in range(1, (ws.max_row or 1) + 1):
        cell = ws.cell(r, col)
        if cell.value is None:
            continue
        max_len = max(max_len, len(_display_text(cell)))
    ws.column_dimensions[get_column_letter(col)].width = min(max(max_len + 2, 8.43), 255)


def _autofit_row(ws, row: int) -> None:
    max_h = _DEFAULT_ROW_HEIGHT
    for c in range(1, (ws.max_column or 1) + 1):
        cell = ws.cell(row, c)
        if cell.value is None:
            continue
        lines = str(_display_text(cell)).count("\n") + 1
        size = cell.font.size or 11
        max_h = max(max_h, lines * float(size) * 1.3)
    ws.row_dimensions[row].height = max_h


def _parse_number(v):
    if v is None or v == "":
        return None
    if isinstance(v, bool):
        return None
    if _is_number(v):
        return v
    s = str(v).strip().replace(",", "")
    if not s:
        return None
    if s.endswith("%"):
        try:
            return float(s[:-1]) / 100.0
        except ValueError:
            return None
    try:
        if "." not in s and "e" not in s.lower():
            return int(s)
        return float(s)
    except ValueError:
        return None


def _numeric_step(vals: list):
    if not vals or not all(_is_number(v) for v in vals):
        return None
    if len(vals) == 1:
        return None
    step = vals[1] - vals[0]
    for i, v in enumerate(vals):
        if abs((v - vals[0]) - step * i) > 1e-9:
            return None
    return step


def _date_step(vals: list):
    if not vals or not all(isinstance(v, (date, datetime)) for v in vals):
        return None
    if len(vals) == 1:
        return None
    step = vals[1] - vals[0]
    for i, v in enumerate(vals):
        if v - vals[0] != step * i:
            return None
    return step


class Range:
    @staticmethod
    def get_range_data(range_obj, use_text: bool = False) -> Any:
        """
        获取区域数据

        Args:
            range_obj: Range 对象
            use_text: 是否返回文本，默认为 False
        """
        try:
            getter = _display_text if use_text else _cell_value
            rows = range_obj.max_row - range_obj.min_row + 1
            cols = range_obj.max_col - range_obj.min_col + 1
            if rows == 1 and cols == 1:
                return getter(range_obj.ws.cell(range_obj.min_row, range_obj.min_col))
            return [
                [getter(range_obj.ws.cell(r, c)) for c in range(range_obj.min_col, range_obj.max_col + 1)]
                for r in range(range_obj.min_row, range_obj.max_row + 1)
            ]
        except Exception as e:
            raise ValueError(f"获取区域数据失败: {e}")

    @staticmethod
    def get_range_color(range_obj) -> tuple[int, int, int]:
        """
        获取单元格区域的背景颜色（RGB格式）

        Args:
            range_obj: Range 对象

        Returns:
            Tuple[int, int, int]: RGB颜色元组 (r, g, b)，每个值范围 0-255
        """
        try:
            cell = range_obj.ws.cell(range_obj.min_row, range_obj.min_col)
            fill = cell.fill
            if fill is None or not fill.patternType:
                return 255, 255, 255
            color = getattr(fill, "fgColor", None) or getattr(fill, "start_color", None)
            return _color_to_rgb(color)
        except Exception as e:
            raise ValueError(f"获取单元格颜色失败: {e}")

    @staticmethod
    def get_range_size(range_obj) -> tuple[int, int, int, int]:
        """
        获取区域的位置
        """
        ws = range_obj.ws
        left = sum(_col_width(ws, c) * _COL_WIDTH_TO_PT for c in range(1, range_obj.min_col))
        top = sum(_row_height(ws, r) for r in range(1, range_obj.min_row))
        width = sum(_col_width(ws, c) * _COL_WIDTH_TO_PT for c in range(range_obj.min_col, range_obj.max_col + 1))
        height = sum(_row_height(ws, r) for r in range(range_obj.min_row, range_obj.max_row + 1))
        return int(left), int(top), int(width), int(height)

    @staticmethod
    def set_range_data(range_obj, value: Any):
        """
        设置区域数据

        Args:
            range_obj: Range 对象
            value: 要设置的值
        """
        try:
            ws = range_obj.ws
            r0, c0 = range_obj.min_row, range_obj.min_col
            r1, c1 = range_obj.max_row, range_obj.max_col
            if isinstance(value, (list, tuple)):
                if value and isinstance(value[0], (list, tuple)):
                    for i, row in enumerate(value):
                        for j, val in enumerate(row):
                            ws.cell(r0 + i, c0 + j).value = val
                else:
                    for j, val in enumerate(value):
                        ws.cell(r0, c0 + j).value = val
            else:
                for r in range(r0, r1 + 1):
                    for c in range(c0, c1 + 1):
                        ws.cell(r, c).value = value
            _mark_dirty(ws)
        except Exception as e:
            raise ValueError(f"设置区域数据失败: {e}")

    @staticmethod
    def set_range_type(
        range_obj,
        col_width: Optional[str] = None,
        bg_color: Optional[tuple[int, int, int]] = None,
        font_color: Optional[tuple[int, int, int]] = None,
        font_type: FontType = FontType.NO_CHANGE,
        font_name: FontNameType = FontNameType.NO_CHANGE,
        font_size: Optional[int] = None,
        number_format: NumberFormatType = NumberFormatType.NO_CHANGE,
        number_format_other: str = "",
        horizontal_align: HorizontalAlign = HorizontalAlign.NO_CHANGE,
        vertical_align: VerticalAlign = VerticalAlign.NO_CHANGE,
        wrap_text: bool = True,
        design_type: ReadRangeType = ReadRangeType.CELL,
        auto_row_height: bool = False,
        auto_column_width: bool = False,
    ):
        """
        设置区域格式
        """
        ws = range_obj.ws
        if col_width:
            width = float(col_width)
            for c in range(range_obj.min_col, range_obj.max_col + 1):
                ws.column_dimensions[get_column_letter(c)].width = width

        font_kwargs: dict[str, Any] = {}
        if font_color:
            rgb = tuple(font_color)
            font_kwargs["color"] = _rgb_hex((int(rgb[0]), int(rgb[1]), int(rgb[2])))
        else:
            font_kwargs["color"] = _rgb_hex((0, 0, 0))
        if font_name != FontNameType.NO_CHANGE:
            font_kwargs["name"] = font_name.value
        if font_size:
            font_kwargs["size"] = font_size
        if font_type == FontType.BOLD:
            font_kwargs["bold"] = True
            font_kwargs["italic"] = False
        elif font_type == FontType.ITALIC:
            font_kwargs["bold"] = False
            font_kwargs["italic"] = True
        elif font_type == FontType.BOLD_ITALIC:
            font_kwargs["bold"] = True
            font_kwargs["italic"] = True
        elif font_type == FontType.NORMAL:
            font_kwargs["bold"] = False
            font_kwargs["italic"] = False

        fill = None
        if bg_color:
            rgb = tuple(bg_color)
            hex_c = _rgb_hex((int(rgb[0]), int(rgb[1]), int(rgb[2])))
            fill = PatternFill(fill_type="solid", fgColor=hex_c, bgColor=hex_c)
        else:
            fill = PatternFill()

        h_align = None
        if horizontal_align != HorizontalAlign.NO_CHANGE:
            h_align = _H_ALIGN.get(horizontal_align.value)
        v_align = None
        if vertical_align != HorizontalAlign.NO_CHANGE:
            v_align = _V_ALIGN.get(vertical_align.value)

        fmt = None
        if number_format != NumberFormatType.NO_CHANGE:
            fmt = number_format_other if number_format == NumberFormatType.CUSTOM else number_format.value

        wrap = True if wrap_text is True else False
        for r in range(range_obj.min_row, range_obj.max_row + 1):
            for c in range(range_obj.min_col, range_obj.max_col + 1):
                cell = ws.cell(r, c)
                font = copy(cell.font)
                for k, v in font_kwargs.items():
                    setattr(font, k, v)
                cell.font = font
                cell.fill = fill
                align = copy(cell.alignment)
                align.wrap_text = wrap
                if h_align is not None:
                    align.horizontal = h_align
                if v_align is not None:
                    align.vertical = v_align
                cell.alignment = align
                if fmt is not None:
                    cell.number_format = fmt

        if design_type == ReadRangeType.ROW and auto_row_height:
            for r in range(range_obj.min_row, range_obj.max_row + 1):
                _autofit_row(ws, r)
        if design_type == ReadRangeType.COLUMN and auto_column_width:
            for c in range(range_obj.min_col, range_obj.max_col + 1):
                _autofit_column(ws, c)
        _mark_dirty(ws)

    @staticmethod
    def delete_range(range_obj, direction: str = "") -> None:
        """
        删除指定区域，可选择左移或上移

        Args:
            range_obj: Range 对象
            direction: 移动方向，RIGHT_MOVE_LEFT(右侧单元格左移)，LOWER_MOVE_UP(下方单元格上移)
        """
        try:
            ws = range_obj.ws
            n_rows = range_obj.max_row - range_obj.min_row + 1
            n_cols = range_obj.max_col - range_obj.min_col + 1
            if direction == "right_move_left":
                _delete_shift(range_obj, 0, n_cols)
            elif direction == "lower_move_up":
                _delete_shift(range_obj, n_rows, 0)
            else:
                used_rows = max(ws.max_row or 1, 1)
                used_cols = max(ws.max_column or 1, 1)
                full_cols = range_obj.min_row == 1 and range_obj.max_row >= used_rows
                full_rows = range_obj.min_col == 1 and range_obj.max_col >= used_cols
                if full_rows and n_cols >= used_cols:
                    ws.delete_rows(range_obj.min_row, n_rows)
                elif full_cols and n_rows >= used_rows:
                    ws.delete_cols(range_obj.min_col, n_cols)
                elif n_cols >= n_rows:
                    _delete_shift(range_obj, n_rows, 0)
                else:
                    _delete_shift(range_obj, 0, n_cols)
            _mark_dirty(ws)
        except Exception as e:
            raise ValueError(f"删除区域失败: {e}")

    @staticmethod
    def clear_range(range_obj, clear_type: str = ""):
        """
        清理单元格区域内容、格式或全部

        Args:
            range_obj: Range 对象
            clear_type: 清理类型
        """
        try:
            if clear_type == ClearType.CONTENT.value:
                for cell in range_obj.cells():
                    cell.value = None
            elif clear_type == ClearType.STYLE.value:
                for cell in range_obj.cells():
                    _clear_formats(cell)
            elif clear_type == ClearType.ALL.value:
                for cell in range_obj.cells():
                    _clear_all(cell)
            else:
                raise ValueError(f"不支持的清理类型: {clear_type}")
            _mark_dirty(range_obj.ws)
        except Exception as e:
            raise ValueError(f"清理区域失败: {e}")

    @staticmethod
    def copy_range(
        range_obj,
    ):
        """
        拷贝单元格区域
        Args:
            range_obj: Range 对象
        """
        book = _get_book(range_obj.ws)
        clip = _clip_from_range(range_obj)
        book.clipboard = clip
        _system_copy(_to_tsv(_grid_from_range(range_obj)))

    @staticmethod
    def paste_range(
        range_obj,
        paste_type: str = "",
        skip_blanks=False,
        transpose=False,
    ):
        """
        粘贴区域的内容，支持多种粘贴方式

        Args:
            range_obj: Range 对象，粘贴起始区域
            paste_type: 粘贴类型
            skip_blanks: 跳过空白单元格
            transpose: 是否转置粘贴
        """
        if paste_type not in _PASTE_TYPES:
            raise ValueError(f"不支持的粘贴类型: {paste_type}")
        try:
            book = _get_book(range_obj.ws)
            clip = book.clipboard
            if not clip:
                clip = _clip_from_tsv(_system_paste())
            if not clip or not clip.get("cells"):
                raise ValueError("粘贴失败: 剪贴板为空")
            if transpose:
                clip = _transpose_clip(clip)
            _paste_clip(range_obj, clip, paste_type, bool(skip_blanks))
            _mark_dirty(range_obj.ws)
        except ValueError:
            raise
        except Exception as e:
            raise ValueError(f"区域粘贴失败: {e}")

    @staticmethod
    def insert_range(range_obj, axis: str = "row"):
        """
        插入整行或整列

        Args:
            range_obj: Range 对象
            axis: "row" 表示插入行, "column" 表示插入列
        """
        ws = range_obj.ws
        if axis == "row":
            ws.insert_rows(range_obj.min_row, range_obj.max_row - range_obj.min_row + 1)
        elif axis == "column":
            ws.insert_cols(range_obj.min_col, range_obj.max_col - range_obj.min_col + 1)
        else:
            raise ValueError(f"不支持的axis参数: {axis}")
        _mark_dirty(ws)

    @staticmethod
    def merge_range(
        range_obj,
        job_type: str,
    ):
        """
        合并或拆分单元格区域

        Args:
            range_obj: Range 对象
            job_type: 操作类型，MERGE 表示合并，SPLIT 表示拆分
        """
        try:
            addr = range_obj.address
            if job_type == "merge":
                range_obj.ws.merge_cells(addr)
            else:
                range_obj.ws.unmerge_cells(addr)
            _mark_dirty(range_obj.ws)
        except Exception as e:
            raise ValueError(f"合并/拆分单元格失败: {e}")

    @staticmethod
    def autofill_range(range_obj, target_range):
        """
        区域自动填充

        Args:
            range_obj: 起始 Range 对象（要自动填充的单元格/区域）
            target_range: 目标填充范围（Range 对象）
        """
        try:
            src, tgt = range_obj, target_range
            if tgt.ws is not src.ws:
                raise ValueError("自动填充失败: 源与目标不在同一工作表")
            contained = (
                tgt.min_row <= src.min_row <= src.max_row <= tgt.max_row
                and tgt.min_col <= src.min_col <= src.max_col <= tgt.max_col
            )
            if not contained:
                raise ValueError("自动填充失败: 目标区域必须包含源区域")
            ws = src.ws
            for c in range(src.min_col, src.max_col + 1):
                seeds = [ws.cell(r, c) for r in range(src.min_row, src.max_row + 1)]
                seed_vals = [s.value for s in seeds]
                step = _numeric_step(seed_vals)
                dstep = _date_step(seed_vals)
                for dest_r in range(src.max_row + 1, tgt.max_row + 1):
                    dest = ws.cell(dest_r, c)
                    k = dest_r - src.min_row
                    if step is not None:
                        dest.value = seed_vals[0] + step * k
                    elif dstep is not None:
                        dest.value = seed_vals[0] + dstep * k
                    else:
                        src_cell = seeds[k % len(seeds)]
                        dest.value = _shift_formula(src_cell.value, src_cell.coordinate, dest.coordinate)
                        dest.number_format = src_cell.number_format
            for r in range(src.min_row, src.max_row + 1):
                seeds = [ws.cell(r, c) for c in range(src.min_col, src.max_col + 1)]
                seed_vals = [s.value for s in seeds]
                step = _numeric_step(seed_vals)
                dstep = _date_step(seed_vals)
                for dest_c in range(src.max_col + 1, tgt.max_col + 1):
                    dest = ws.cell(r, dest_c)
                    k = dest_c - src.min_col
                    if step is not None:
                        dest.value = seed_vals[0] + step * k
                    elif dstep is not None:
                        dest.value = seed_vals[0] + dstep * k
                    else:
                        src_cell = seeds[k % len(seeds)]
                        dest.value = _shift_formula(src_cell.value, src_cell.coordinate, dest.coordinate)
                        dest.number_format = src_cell.number_format
            _mark_dirty(ws)
        except ValueError:
            raise
        except Exception as e:
            raise ValueError(f"区域自动填充失败: {e}")

    @staticmethod
    def set_row_height(range_obj, set_type: SetType, height_float: float):
        """
        设置行高
        """
        ws = range_obj.ws
        if set_type == SetType.VALUE:
            for r in range(range_obj.min_row, range_obj.max_row + 1):
                ws.row_dimensions[r].height = height_float
        elif set_type == SetType.AUTO:
            for r in range(range_obj.min_row, range_obj.max_row + 1):
                _autofit_row(ws, r)
        _mark_dirty(ws)

    @staticmethod
    def set_column_width(range_obj, set_type: SetType, width_float: float):
        """
        设置列宽
        """
        ws = range_obj.ws
        if set_type == SetType.VALUE:
            for c in range(range_obj.min_col, range_obj.max_col + 1):
                ws.column_dimensions[get_column_letter(c)].width = width_float
        elif set_type == SetType.AUTO:
            for c in range(range_obj.min_col, range_obj.max_col + 1):
                _autofit_column(ws, c)
        _mark_dirty(ws)

    @staticmethod
    def convert_text_to_number(range_obj, temp_range):
        """
        将范围内的文本格式转换为数值格式

        Args:
            range_obj: Range 对象（要转换的范围）
            temp_range: 临时单元格 Range 对象（openpyxl 不计算 VALUE，忽略）
        """
        for cell in range_obj.cells():
            if cell.value in ("", None):
                continue
            num = _parse_number(cell.value)
            if num is None:
                continue
            cell.number_format = "G/通用格式"
            cell.value = num
        _mark_dirty(range_obj.ws)

    @staticmethod
    def convert_number_to_text(range_obj):
        """
        将范围内的数值格式转换为文本格式
        """
        for cell in range_obj.cells():
            if not isinstance(cell.value, str):
                text = _display_text(cell)
                cell.number_format = "@"
                cell.value = text
        _mark_dirty(range_obj.ws)

    @staticmethod
    def add_comment(range_obj, comment_text: str):
        """
        为范围添加批注
        """
        cell = range_obj.ws.cell(range_obj.min_row, range_obj.min_col)
        cell.comment = Comment(comment_text, "AstronRPA")
        _mark_dirty(range_obj.ws)

    @staticmethod
    def delete_comment(range_obj):
        """
        删除范围的批注
        """
        cell = range_obj.ws.cell(range_obj.min_row, range_obj.min_col)
        if cell.comment:
            for c in range_obj.cells():
                c.comment = None
            _mark_dirty(range_obj.ws)
        else:
            raise ValueError("不存在批注")

    @staticmethod
    def search_and_replace(
        range_obj,
        find_str: str,
        replace_str: str = "",
        exact_match: bool = False,
        case_flag: bool = False,
        match_all: bool = True,
    ) -> list:
        """
        在范围内搜索并可选地替换文本
        """
        import re

        positions = []
        for cell in range_obj.cells():
            if cell.value is None:
                continue
            cell_value = str(cell.value)
            hay = cell_value if case_flag else cell_value.lower()
            needle = find_str if case_flag else find_str.lower()
            matched = hay == needle if exact_match else needle in hay
            if not matched:
                continue
            positions.append({"row": str(cell.row), "col": column_number_to_letter(cell.column)})
            if replace_str:
                if case_flag:
                    cell.value = cell_value.replace(find_str, replace_str)
                else:
                    cell.value = re.sub(re.escape(find_str), lambda _: replace_str, cell_value, flags=re.IGNORECASE)
            if not match_all:
                break
        if replace_str and positions:
            _mark_dirty(range_obj.ws)
        return positions
