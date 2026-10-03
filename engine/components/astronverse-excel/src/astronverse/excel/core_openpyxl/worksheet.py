from copy import copy

from astronverse.excel import CopySheetLocationType
from astronverse.excel.core_openpyxl import OpenpyxlBook, OpenpyxlRange, book_of, cell_address
from astronverse.excel.excel_obj import ExcelObj
from openpyxl.utils import get_column_letter


def _wb(excel_obj: ExcelObj):
    obj = excel_obj.obj
    return obj.wb if isinstance(obj, OpenpyxlBook) else obj


def _mark_dirty(ws_or_wb) -> None:
    book_of(ws_or_wb).dirty = True


def _resolve_sheet(wb, ref):
    try:
        return wb[ref]
    except (KeyError, TypeError):
        pass
    try:
        idx = int(ref)
        sheets = wb.worksheets
        if 1 <= idx <= len(sheets):
            return sheets[idx - 1]
    except (TypeError, ValueError):
        pass
    raise ValueError(f"工作表'{ref}'不存在")


def _unique_title(wb, title: str) -> str:
    if title not in wb.sheetnames:
        return title
    n = 1
    while f"{title} ({n})" in wb.sheetnames:
        n += 1
    return f"{title} ({n})"


def _copy_sheet_contents(src, dest):
    for row in src.iter_rows():
        for cell in row:
            d = dest.cell(cell.row, cell.column, cell.value)
            if cell.has_style:
                d.font = copy(cell.font)
                d.fill = copy(cell.fill)
                d.border = copy(cell.border)
                d.alignment = copy(cell.alignment)
                d.number_format = cell.number_format
                d.protection = copy(cell.protection)
            if cell.comment:
                d.comment = copy(cell.comment)
            if cell.hyperlink:
                d.hyperlink = copy(cell.hyperlink)
    for merged in src.merged_cells.ranges:
        dest.merge_cells(str(merged))
    for letter, dim in src.column_dimensions.items():
        dest.column_dimensions[letter].width = dim.width
        dest.column_dimensions[letter].hidden = dim.hidden
    for idx, dim in src.row_dimensions.items():
        dest.row_dimensions[idx].height = dim.height
        dest.row_dimensions[idx].hidden = dim.hidden


def _move_sheet_to_index(wb, ws, index: int):
    sheets = wb._sheets
    sheets.remove(ws)
    index = max(0, min(index, len(sheets)))
    sheets.insert(index, ws)


class Worksheet:
    @staticmethod
    def get_worksheet(excel_obj: ExcelObj, sheet_name: str = "", default: int = 0) -> object:
        wb = _wb(excel_obj)
        if not sheet_name:
            if default == 0:
                return Worksheet.get_active_worksheet(excel_obj)
            else:
                sheet_name = default
        return _resolve_sheet(wb, sheet_name)

    @staticmethod
    def get_all_worksheets(excel_obj: ExcelObj) -> list[object]:
        return list(_wb(excel_obj).worksheets)

    @staticmethod
    def get_all_worksheet_names(excel_obj: ExcelObj) -> list[str]:
        return list(_wb(excel_obj).sheetnames)

    @staticmethod
    def get_active_worksheet(excel_obj: ExcelObj) -> object:
        return _wb(excel_obj).active

    @staticmethod
    def add_worksheet(excel_obj: ExcelObj, sheet_name: str, before=None, after=None):
        wb = _wb(excel_obj)
        if before is not None:
            ref = _resolve_sheet(wb, before)
            idx = wb.worksheets.index(ref)
        elif after is not None:
            ref = _resolve_sheet(wb, after)
            idx = wb.worksheets.index(ref) + 1
        else:
            idx = len(wb.worksheets)
        new_sheet = wb.create_sheet(title=sheet_name, index=idx)
        _mark_dirty(wb)
        return new_sheet

    @staticmethod
    def move_worksheet(worksheet, before=None, after=None):
        wb = worksheet.parent
        if before is not None:
            ref = _resolve_sheet(wb, before)
            idx = wb.worksheets.index(ref)
            if wb.worksheets.index(worksheet) < idx:
                idx -= 1
        elif after is not None:
            ref = _resolve_sheet(wb, after)
            idx = wb.worksheets.index(ref) + 1
            if wb.worksheets.index(worksheet) < idx:
                idx -= 1
        else:
            idx = len(wb.worksheets) - 1
        _move_sheet_to_index(wb, worksheet, idx)
        _mark_dirty(wb)

    @staticmethod
    def get_worksheet_name(worksheet) -> str:
        return worksheet.title

    @staticmethod
    def rename_worksheet(worksheet, new_name: str):
        worksheet.title = new_name
        _mark_dirty(worksheet)

    @staticmethod
    def delete_worksheet(worksheet):
        wb = worksheet.parent
        if len(wb.worksheets) <= 1:
            raise ValueError("不能删除唯一的工作表")
        wb.remove(worksheet)
        _mark_dirty(wb)

    @staticmethod
    def copy_worksheet(
        worksheet, excel, location: CopySheetLocationType = CopySheetLocationType.LAST, is_same_workbook=False
    ):
        src_wb = worksheet.parent
        dest_wb = _wb(excel)
        same = is_same_workbook or src_wb is dest_wb
        if same:
            new_ws = dest_wb.copy_worksheet(worksheet)
            src_idx = dest_wb.worksheets.index(worksheet)
            if location == CopySheetLocationType.BEFORE:
                _move_sheet_to_index(dest_wb, new_ws, src_idx)
            elif location == CopySheetLocationType.AFTER:
                _move_sheet_to_index(dest_wb, new_ws, src_idx + 1)
            elif location == CopySheetLocationType.FIRST:
                _move_sheet_to_index(dest_wb, new_ws, 0)
            dest_wb.active = new_ws
        else:
            title = _unique_title(dest_wb, worksheet.title)
            if location == CopySheetLocationType.BEFORE:
                idx = dest_wb.worksheets.index(dest_wb.active)
            elif location == CopySheetLocationType.AFTER:
                idx = dest_wb.worksheets.index(dest_wb.active) + 1
            elif location == CopySheetLocationType.FIRST:
                idx = 0
            else:
                idx = len(dest_wb.worksheets)
            new_ws = dest_wb.create_sheet(title=title, index=idx)
            _copy_sheet_contents(worksheet, new_ws)
            dest_wb.active = new_ws
        _mark_dirty(dest_wb)
        return new_ws

    @staticmethod
    def get_worksheet_used_range(worksheet):
        """按实际非空单元格计算已用区域，裁掉尾部空行/空列。"""
        min_row = min_col = max_row = max_col = None
        for row in worksheet.iter_rows(min_row=1, max_row=worksheet.max_row, max_col=worksheet.max_column):
            for cell in row:
                val = cell.value
                if val is None or val == "":
                    continue
                r, c = cell.row, cell.column
                if min_row is None or r < min_row:
                    min_row = r
                if max_row is None or r > max_row:
                    max_row = r
                if min_col is None or c < min_col:
                    min_col = c
                if max_col is None or c > max_col:
                    max_col = c
        if min_row is None:
            return 1, 1, 0, 0, "A1"
        start = cell_address(min_row, min_col)
        end = cell_address(max_row, max_col)
        address = start if start == end else f"{start}:{end}"
        return min_row, min_col, max_row, max_col, address

    @staticmethod
    def get_cell(worksheet, row: int, col: int) -> object:
        try:
            return OpenpyxlRange(worksheet, row, col, row, col)
        except Exception as e:
            raise ValueError(f"获取单元格({row}, {col})失败: {e}")

    @staticmethod
    def get_range(worksheet, cell: str) -> object:
        try:
            return OpenpyxlRange.from_address(worksheet, cell)
        except Exception as e:
            raise ValueError(f"获取区域 '{cell}' 失败: {e}")

    @staticmethod
    def get_rows(worksheet, rows) -> object:
        try:
            if isinstance(rows, int):
                max_col = worksheet.max_column or 1
                return OpenpyxlRange(worksheet, rows, 1, rows, max_col)
            return OpenpyxlRange.from_address(worksheet, str(rows) if ":" in str(rows) else f"{rows}:{rows}")
        except Exception as e:
            raise ValueError(f"获取区域 '{rows}' 失败: {e}")

    @staticmethod
    def get_columns(worksheet, columns) -> object:
        try:
            if isinstance(columns, int):
                max_row = worksheet.max_row or 1
                return OpenpyxlRange(worksheet, 1, columns, max_row, columns)
            text = str(columns)
            if text.isdigit():
                idx = int(text)
                max_row = worksheet.max_row or 1
                return OpenpyxlRange(worksheet, 1, idx, max_row, idx)
            if ":" not in text:
                letter = get_column_letter(int(text)) if text.isdigit() else text
                text = f"{letter}:{letter}"
            return OpenpyxlRange.from_address(worksheet, text)
        except Exception as e:
            raise ValueError(f"获取区域 '{columns}' 失败: {e}")

    @staticmethod
    def get_range_from_cells(worksheet, start_cell, end_cell) -> object:
        try:
            min_row = min(start_cell.min_row, end_cell.min_row)
            min_col = min(start_cell.min_col, end_cell.min_col)
            max_row = max(start_cell.max_row, end_cell.max_row)
            max_col = max(start_cell.max_col, end_cell.max_col)
            return OpenpyxlRange(worksheet, min_row, min_col, max_row, max_col)
        except Exception as e:
            raise ValueError(f"通过 Range 对象获取区域失败: {e}")

    @staticmethod
    def insert_picture(worksheet, image_path, pic_left=0, pic_top=0, pic_height=300, pic_width=400, pic_scale=1.0):
        from openpyxl.drawing.image import Image as XLImage
        from PIL import Image as PILImage

        picture = XLImage(image_path)
        if pic_scale != 1.0:
            with PILImage.open(image_path) as image:
                width, height = image.size
            picture.width = width * pic_scale
            picture.height = height * pic_scale
        else:
            picture.width = pic_width
            picture.height = pic_height
        col = max(1, int(pic_left / 64) + 1)
        row = max(1, int(pic_top / 20) + 1)
        picture.anchor = cell_address(row, col)
        worksheet.add_image(picture)
        _mark_dirty(worksheet)
        return picture

    @staticmethod
    def delete_all_comments(worksheet):
        found = False
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.comment:
                    cell.comment = None
                    found = True
        if found:
            _mark_dirty(worksheet)
        else:
            raise ValueError("不存在批注")
