import builtins
from dataclasses import dataclass
from typing import Optional
from weakref import WeakKeyDictionary

from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

_BOOKS: WeakKeyDictionary = WeakKeyDictionary()


@dataclass
class OpenpyxlBook:
    wb: Workbook
    path: str = ""
    name: str = ""
    password: str = ""
    visible: bool = False
    dirty: bool = False
    clipboard: dict | None = None


def register_book(book: OpenpyxlBook) -> OpenpyxlBook:
    _BOOKS[book.wb] = book
    book.wb._openpyxl_book = book
    return book


def unregister_book(book: OpenpyxlBook) -> None:
    _BOOKS.pop(book.wb, None)
    if getattr(book.wb, "_openpyxl_book", None) is book:
        book.wb._openpyxl_book = None


def book_of(ws_or_wb) -> OpenpyxlBook:
    wb = ws_or_wb if isinstance(ws_or_wb, Workbook) else ws_or_wb.parent
    book = _BOOKS.get(wb) or getattr(wb, "_openpyxl_book", None)
    if book is None:
        raise ValueError("未找到对应的工作簿")
    return book


def cell_address(row: int, col: int) -> str:
    return f"{get_column_letter(col)}{row}"


def col_letter(col: int) -> str:
    return get_column_letter(col)


def col_index(letter: str) -> int:
    return column_index_from_string(letter)


def parse_address(addr: str) -> tuple[Optional[int], Optional[int], Optional[int], Optional[int]]:
    """返回 (min_col, min_row, max_col, max_row)，整行/整列缺失的边界为 None。"""
    return range_boundaries(addr.replace("$", "").strip())


@dataclass
class OpenpyxlRange:
    ws: Worksheet
    min_row: int
    min_col: int
    max_row: int
    max_col: int

    @property
    def address(self) -> str:
        start = cell_address(self.min_row, self.min_col)
        end = cell_address(self.max_row, self.max_col)
        return start if start == end else f"{start}:{end}"

    def cells(self):
        for row in builtins.range(self.min_row, self.max_row + 1):
            for col in builtins.range(self.min_col, self.max_col + 1):
                yield self.ws.cell(row, col)

    @classmethod
    def from_address(cls, ws: Worksheet, addr: str) -> "OpenpyxlRange":
        min_col, min_row, max_col, max_row = parse_address(addr)
        if min_row is None:
            min_row = 1
        if max_row is None:
            max_row = ws.max_row or 1
        if min_col is None:
            min_col = 1
        if max_col is None:
            max_col = ws.max_column or 1
        return cls(ws, min_row, min_col, max_row, max_col)
