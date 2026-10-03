import os
import tempfile
from typing import Optional

import openpyxl
from astronverse.excel import ApplicationType
from astronverse.excel.core_openpyxl import OpenpyxlBook, register_book, unregister_book
from astronverse.excel.excel_obj import ExcelObj

_APP = None


def get_default_excel_application():
    """macOS 无 COM/注册表，固定返回 Excel。"""
    return ApplicationType.EXCEL


class OpenpyxlApp:
    """文件型 Excel 应用：跟踪已打开工作簿。visible=True 时不自动拉起任何桌面应用。"""

    def __init__(self, visible: bool = False):
        self.visible = visible
        self.books: list[OpenpyxlBook] = []


def _app() -> Optional[OpenpyxlApp]:
    return _APP


def _ensure_app(visible_flag: Optional[bool] = None) -> OpenpyxlApp:
    global _APP
    if _APP is None:
        _APP = OpenpyxlApp(visible=bool(visible_flag))
    elif visible_flag is not None:
        _APP.visible = visible_flag
    return _APP


def _is_xls(path: str) -> bool:
    return os.path.splitext(path)[1].lower() == ".xls"


def _decrypt_to_temp(path: str, password: str) -> Optional[str]:
    try:
        import msoffcrypto
    except ImportError:
        return None
    with open(path, "rb") as inf:
        office = msoffcrypto.OfficeFile(inf)
        if not office.is_encrypted():
            return None
        if not password:
            raise Exception("文件已加密，openpyxl 无法解密，请提供密码")
        fd, tmp_path = tempfile.mkstemp(suffix=os.path.splitext(path)[1] or ".xlsx")
        os.close(fd)
        inf.seek(0)
        with open(tmp_path, "wb") as out:
            office.load_key(password=password)
            office.decrypt(out)
        return tmp_path


def _load_workbook(path: str, password: str = ""):
    if _is_xls(path):
        raise Exception("不支持 .xls 格式，请先将文件转换为 .xlsx 后再打开")
    tmp_path = None
    load_path = path
    try:
        try:
            tmp_path = _decrypt_to_temp(path, password)
            if tmp_path:
                load_path = tmp_path
        except Exception as e:
            if "已加密" in str(e) or "password" in str(e).lower() or "encrypt" in str(e).lower():
                raise
            tmp_path = None
        keep_vba = path.lower().endswith(".xlsm")
        try:
            return openpyxl.load_workbook(load_path, keep_vba=keep_vba, data_only=False)
        except Exception as e:
            if password:
                raise Exception(
                    f"无法打开受密码保护的 Excel 文件（openpyxl 无法直接解密，请安装 msoffcrypto-tool）: {e}"
                )
            msg = str(e).lower()
            if "zip" in msg or "encrypt" in msg or "password" in msg:
                raise Exception(f"无法打开 Excel 文件（若文件受密码保护，请安装 msoffcrypto-tool 并提供正确密码）: {e}")
            raise
    finally:
        if tmp_path:
            try:
                os.unlink(tmp_path)
            except OSError:
                pass


def _book_name(path: str, fallback: str = "Book1.xlsx") -> str:
    return os.path.basename(path) if path else fallback


class Application:
    @staticmethod
    def init_app(
        default_application: ApplicationType = ApplicationType.DEFAULT,
        visible_flag: Optional[bool] = None,
        retry: int = 0,
        retry_delay: float = 0.5,
        prefer_existing: bool = True,
    ) -> object:
        """初始化文件型 Excel 应用（不启动桌面程序）。"""
        return _ensure_app(visible_flag)

    @staticmethod
    def quit_app(default_application: ApplicationType = ApplicationType.DEFAULT, save_changes: bool = False):
        """关闭全部已打开工作簿。visible 时也不自动打开 Excel/Numbers。"""
        global _APP
        app = _APP
        if app is None:
            return
        for book in list(app.books):
            try:
                if save_changes and book.path:
                    book.wb.save(book.path)
                    book.dirty = False
                book.wb.close()
            except Exception:
                pass
            unregister_book(book)
        app.books.clear()
        _APP = None

    @staticmethod
    def create_workbook(application, file_path: str = "", password: str = "") -> ExcelObj:
        """创建新工作簿"""
        wb = openpyxl.Workbook()
        if wb.active is not None:
            wb.active.title = "Sheet1"
        path = os.path.abspath(file_path) if file_path else ""
        book = register_book(
            OpenpyxlBook(
                wb=wb,
                path=path,
                name=_book_name(path),
                password=password or "",
                visible=bool(getattr(application, "visible", False)),
                dirty=not bool(path),
            )
        )
        application.books.append(book)
        if path:
            wb.save(path)
            book.dirty = False
        return ExcelObj(obj=book, path=path)

    @staticmethod
    def open_workbook(application, file_path: str, password: str = "", update_links: bool = True) -> ExcelObj:
        """打开工作簿（data_only=False；.xls 不支持）。"""
        path = os.path.abspath(file_path)
        wb = _load_workbook(path, password=password)
        book = register_book(
            OpenpyxlBook(
                wb=wb,
                path=path,
                name=_book_name(path),
                password=password or "",
                visible=bool(getattr(application, "visible", False)),
                dirty=False,
            )
        )
        application.books.append(book)
        return ExcelObj(obj=book, path=path)

    @staticmethod
    def get_existing_workbook(application, match_name: str) -> Optional[ExcelObj]:
        """获取已打开的工作簿"""
        books = getattr(application, "books", None) or []
        for book in reversed(books):
            if match_name in book.name:
                return ExcelObj(obj=book, path=book.path)
        return None

    @staticmethod
    def save_workbook(excel_obj: ExcelObj, file_path: str = "", password: str = ""):
        book: OpenpyxlBook = excel_obj.obj
        target = os.path.abspath(file_path) if file_path else book.path
        if not target:
            raise Exception("未指定保存路径")
        if password:
            book.password = password
        book.wb.save(target)
        book.path = target
        book.name = _book_name(target)
        book.dirty = False
        excel_obj.path = target

    @staticmethod
    def close_workbook(excel_obj: ExcelObj, save_changes: bool = True):
        book: OpenpyxlBook = excel_obj.obj
        if save_changes:
            if book.path:
                book.wb.save(book.path)
                book.dirty = False
            elif book.dirty:
                raise Exception("未指定保存路径")
        try:
            book.wb.close()
        except Exception:
            pass
        unregister_book(book)
        app = _app()
        if app is not None and book in app.books:
            app.books.remove(book)
