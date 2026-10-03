import importlib
import importlib.util
import sys
import types
from datetime import date, datetime
from pathlib import Path
from types import SimpleNamespace

import pytest
from astronverse.excel import (
    ClearType,
    FontNameType,
    FontType,
    HorizontalAlign,
    NumberFormatType,
    ReadRangeType,
    SetType,
    VerticalAlign,
)
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter, range_boundaries

_RANGE_PY = Path(__file__).resolve().parents[1] / "src" / "astronverse" / "excel" / "core_openpyxl" / "range.py"


class StubRange:
    def __init__(self, ws, min_row, min_col, max_row, max_col):
        self.ws = ws
        self.min_row = min_row
        self.min_col = min_col
        self.max_row = max_row
        self.max_col = max_col

    @property
    def address(self) -> str:
        a = f"{get_column_letter(self.min_col)}{self.min_row}"
        b = f"{get_column_letter(self.max_col)}{self.max_row}"
        return a if a == b else f"{a}:{b}"

    def cells(self):
        for r in range(self.min_row, self.max_row + 1):
            for c in range(self.min_col, self.max_col + 1):
                yield self.ws.cell(r, c)

    @classmethod
    def from_address(cls, ws, addr: str):
        min_col, min_row, max_col, max_row = range_boundaries(addr)
        return cls(ws, min_row, min_col, max_row, max_col)


def _load_range_mod():
    try:
        return importlib.import_module("astronverse.excel.core_openpyxl.range")
    except ModuleNotFoundError:
        spec = importlib.util.spec_from_file_location("astronverse.excel.core_openpyxl.range", _RANGE_PY)
        mod = importlib.util.module_from_spec(spec)
        pkg = types.ModuleType("astronverse.excel.core_openpyxl")
        pkg.__path__ = [str(_RANGE_PY.parent)]
        sys.modules.setdefault("astronverse.excel.core_openpyxl", pkg)
        sys.modules[spec.name] = mod
        spec.loader.exec_module(mod)
        return mod


def _try_openpyxl_range():
    try:
        from astronverse.excel.core_openpyxl import OpenpyxlRange

        return OpenpyxlRange
    except Exception:
        return StubRange


class _Clip:
    def __init__(self):
        self.text = ""

    def copy(self, text):
        self.text = text

    def paste(self):
        return self.text


@pytest.fixture
def xl(monkeypatch):
    mod = _load_range_mod()
    Range = mod.Range
    rng_cls = _try_openpyxl_range()
    wb = Workbook()
    ws = wb.active
    book = SimpleNamespace(wb=wb, path="", name="Book", dirty=False, clipboard=None)
    wb._openpyxl_book = book

    def fake_book_of(ws_):
        parent = getattr(ws_, "parent", None)
        return getattr(parent, "_openpyxl_book", book)

    monkeypatch.setattr(mod, "_get_book", fake_book_of)
    clip = _Clip()
    monkeypatch.setattr(mod, "_system_copy", clip.copy)
    monkeypatch.setattr(mod, "_system_paste", clip.paste)

    def rng(addr: str, sheet=None):
        return rng_cls.from_address(sheet or ws, addr)

    return SimpleNamespace(Range=Range, wb=wb, ws=ws, book=book, rng=rng, clip=clip, rng_cls=rng_cls)


def test_get_range_data_scalar_and_grid(xl):
    xl.ws["A1"] = 12
    xl.ws["B1"] = "x"
    xl.ws["A2"] = 3.5
    assert xl.Range.get_range_data(xl.rng("A1")) == 12
    assert xl.Range.get_range_data(xl.rng("A1:B2")) == [[12, "x"], [3.5, None]]


def test_get_range_data_use_text_and_formula(xl):
    xl.ws["A1"] = 12
    xl.ws["A1"].number_format = "0.00"
    xl.ws["B1"] = "=A1+1"
    assert xl.Range.get_range_data(xl.rng("A1"), use_text=True) == "12.00"
    assert xl.Range.get_range_data(xl.rng("B1")) == "=A1+1"
    assert xl.Range.get_range_data(xl.rng("B1"), use_text=True) == "=A1+1"
    xl.ws["C1"] = datetime(2020, 1, 2, 3, 4, 5)
    assert "2020" in xl.Range.get_range_data(xl.rng("C1"), use_text=True)


def test_get_range_color_default_and_rgb(xl):
    assert xl.Range.get_range_color(xl.rng("A1")) == (255, 255, 255)
    xl.ws["A1"].fill = PatternFill(fill_type="solid", fgColor="FF112233")
    assert xl.Range.get_range_color(xl.rng("A1")) == (0x11, 0x22, 0x33)


def test_get_range_size(xl):
    left, top, width, height = xl.Range.get_range_size(xl.rng("A1"))
    assert left == 0
    assert top == 0
    assert width > 0
    assert height > 0
    l2, t2, w2, h2 = xl.Range.get_range_size(xl.rng("B2"))
    assert l2 > left
    assert t2 > top
    _, _, w3, h3 = xl.Range.get_range_size(xl.rng("A1:B2"))
    assert w3 > width
    assert h3 > height


def test_set_range_data_scalar_list_grid(xl):
    xl.Range.set_range_data(xl.rng("A1"), "hi")
    assert xl.ws["A1"].value == "hi"
    assert xl.book.dirty is True
    xl.book.dirty = False
    xl.Range.set_range_data(xl.rng("A1:C1"), [1, 2, 3])
    assert [xl.ws.cell(1, c).value for c in range(1, 4)] == [1, 2, 3]
    xl.Range.set_range_data(xl.rng("A2:B3"), [[4, 5], [6, 7]])
    assert xl.ws["A2"].value == 4
    assert xl.ws["B3"].value == 7
    xl.Range.set_range_data(xl.rng("A4:B4"), 9)
    assert xl.ws["A4"].value == 9
    assert xl.ws["B4"].value == 9


def test_set_range_type_font_fill_align(xl):
    xl.ws["A1"] = 1
    xl.Range.set_range_type(
        xl.rng("A1"),
        col_width="12",
        bg_color=(10, 20, 30),
        font_color=(1, 2, 3),
        font_type=FontType.BOLD,
        font_name=FontNameType.SONGTI,
        font_size=14,
        number_format=NumberFormatType.NUMBER,
        horizontal_align=HorizontalAlign.CENTER,
        vertical_align=VerticalAlign.MIDDLE,
        wrap_text=True,
    )
    cell = xl.ws["A1"]
    assert cell.font.bold is True
    assert cell.font.size == 14
    assert cell.font.name == "宋体"
    assert cell.alignment.horizontal == "center"
    assert cell.alignment.vertical == "center"
    assert cell.alignment.wrap_text is True
    assert cell.number_format == "0.00"
    assert xl.ws.column_dimensions["A"].width == 12
    assert xl.Range.get_range_color(xl.rng("A1")) == (10, 20, 30)
    assert xl.book.dirty is True


def test_set_range_type_autofit(xl):
    xl.ws["A1"] = "abcdefghijklmnop"
    xl.Range.set_range_type(
        xl.rng("A1"),
        design_type=ReadRangeType.COLUMN,
        auto_column_width=True,
        wrap_text=False,
    )
    assert xl.ws.column_dimensions["A"].width >= 8.43
    xl.ws["B1"] = "line1\nline2\nline3"
    xl.Range.set_range_type(
        xl.rng("B1"),
        design_type=ReadRangeType.ROW,
        auto_row_height=True,
        wrap_text=True,
    )
    assert xl.ws.row_dimensions[1].height > 15


def test_delete_range_shift_up(xl):
    xl.ws["A1"] = "a"
    xl.ws["A2"] = "b"
    xl.ws["A3"] = "c"
    xl.Range.delete_range(xl.rng("A1"), direction="lower_move_up")
    assert xl.ws["A1"].value == "b"
    assert xl.ws["A2"].value == "c"
    assert xl.ws["A3"].value is None
    assert xl.book.dirty is True


def test_delete_range_shift_left(xl):
    xl.ws["A1"] = "a"
    xl.ws["B1"] = "b"
    xl.ws["C1"] = "c"
    xl.Range.delete_range(xl.rng("A1"), direction="right_move_left")
    assert xl.ws["A1"].value == "b"
    assert xl.ws["B1"].value == "c"
    assert xl.ws["C1"].value is None


def test_delete_range_whole_row(xl):
    xl.ws["A1"] = 1
    xl.ws["B1"] = 2
    xl.ws["A2"] = 3
    xl.Range.delete_range(xl.rng("A1:B1"))
    assert xl.ws["A1"].value == 3


def test_clear_range_content_style_all(xl):
    cell = xl.ws["A1"]
    cell.value = "x"
    cell.font = Font(bold=True)
    cell.fill = PatternFill(fill_type="solid", fgColor="FF112233")
    xl.Range.clear_range(xl.rng("A1"), ClearType.CONTENT.value)
    assert cell.value is None
    assert cell.font.bold is True
    cell.value = "y"
    xl.Range.clear_range(xl.rng("A1"), ClearType.STYLE.value)
    assert cell.value == "y"
    assert not cell.font.bold
    xl.Range.add_comment(xl.rng("A1"), "c")
    xl.Range.clear_range(xl.rng("A1"), ClearType.ALL.value)
    assert cell.value is None
    assert cell.comment is None
    with pytest.raises(ValueError, match="不支持的清理类型"):
        xl.Range.clear_range(xl.rng("A1"), "nope")


def test_copy_paste_all_and_value(xl):
    xl.ws["A1"] = 11
    xl.ws["A1"].font = Font(bold=True)
    xl.ws["B1"] = 22
    xl.Range.copy_range(xl.rng("A1:B1"))
    assert "11" in xl.clip.text
    assert "22" in xl.clip.text
    xl.Range.paste_range(xl.rng("A3"), paste_type="all")
    assert xl.ws["A3"].value == 11
    assert xl.ws["B3"].value == 22
    assert xl.ws["A3"].font.bold is True
    xl.ws["A4"] = None
    xl.Range.paste_range(xl.rng("A4"), paste_type="paste_value")
    assert xl.ws["A4"].value == 11


def test_paste_format_only_skip_blanks_transpose(xl):
    xl.ws["A1"] = "src"
    xl.ws["A1"].font = Font(italic=True, size=20)
    xl.ws["B1"] = None
    xl.Range.copy_range(xl.rng("A1:B1"))
    xl.ws["A2"] = "keep"
    xl.Range.paste_range(xl.rng("A2"), paste_type="format")
    assert xl.ws["A2"].value == "keep"
    assert xl.ws["A2"].font.italic is True
    xl.ws["C1"] = "old"
    xl.ws["D1"] = "keep-blank"
    xl.Range.paste_range(xl.rng("C1"), paste_type="all", skip_blanks=True)
    assert xl.ws["C1"].value == "src"
    assert xl.ws["D1"].value == "keep-blank"
    xl.ws["A5"] = 1
    xl.ws["B5"] = 2
    xl.Range.copy_range(xl.rng("A5:B5"))
    xl.Range.paste_range(xl.rng("A6"), paste_type="paste_value", transpose=True)
    assert xl.ws["A6"].value == 1
    assert xl.ws["A7"].value == 2


def test_paste_from_system_clipboard(xl):
    xl.book.clipboard = None
    xl.clip.text = "p\tq\nr\ts"
    xl.Range.paste_range(xl.rng("A1"), paste_type="paste_value")
    assert xl.ws["A1"].value == "p"
    assert xl.ws["B2"].value == "s"
    with pytest.raises(ValueError, match="不支持的粘贴类型"):
        xl.Range.paste_range(xl.rng("A1"), paste_type="nope")
    xl.book.clipboard = None
    xl.clip.text = ""
    with pytest.raises(ValueError, match="剪贴板为空"):
        xl.Range.paste_range(xl.rng("A1"), paste_type="all")


def test_insert_range_row_and_column(xl):
    xl.ws["A1"] = "top"
    xl.ws["A2"] = "bot"
    xl.Range.insert_range(xl.rng("A1"), axis="row")
    assert xl.ws["A1"].value is None
    assert xl.ws["A3"].value == "bot"
    xl.ws["B1"] = "x"
    xl.Range.insert_range(xl.rng("A1"), axis="column")
    assert xl.ws["C1"].value == "x"
    with pytest.raises(ValueError, match="axis"):
        xl.Range.insert_range(xl.rng("A1"), axis="diag")


def test_merge_unmerge(xl):
    xl.ws["A1"] = "m"
    xl.Range.merge_range(xl.rng("A1:B2"), "merge")
    assert "A1:B2" in [str(m) for m in xl.ws.merged_cells.ranges]
    xl.Range.merge_range(xl.rng("A1:B2"), "split")
    assert "A1:B2" not in [str(m) for m in xl.ws.merged_cells.ranges]


def test_autofill_numeric_formula_copy(xl):
    xl.ws["A1"] = 1
    xl.ws["A2"] = 3
    xl.Range.autofill_range(xl.rng("A1:A2"), xl.rng("A1:A5"))
    assert [xl.ws.cell(r, 1).value for r in range(1, 6)] == [1, 3, 5, 7, 9]
    xl.ws["B1"] = "=A1"
    xl.Range.autofill_range(xl.rng("B1"), xl.rng("B1:B3"))
    assert xl.ws["B2"].value == "=A2"
    assert xl.ws["B3"].value == "=A3"
    xl.ws["C1"] = "x"
    xl.Range.autofill_range(xl.rng("C1"), xl.rng("C1:C3"))
    assert xl.ws["C2"].value == "x"
    assert xl.ws["C3"].value == "x"
    xl.ws["D1"] = 10
    xl.ws["E1"] = 20
    xl.Range.autofill_range(xl.rng("D1:E1"), xl.rng("D1:G1"))
    assert xl.ws["F1"].value == 30
    assert xl.ws["G1"].value == 40
    xl.ws["A10"] = date(2020, 1, 1)
    xl.ws["A11"] = date(2020, 1, 3)
    xl.Range.autofill_range(xl.rng("A10:A11"), xl.rng("A10:A13"))
    assert xl.ws["A13"].value == date(2020, 1, 7)


def test_set_row_height_and_column_width(xl):
    xl.Range.set_row_height(xl.rng("A1"), SetType.VALUE, 22.5)
    assert xl.ws.row_dimensions[1].height == 22.5
    xl.ws["A2"] = "hello world"
    xl.Range.set_row_height(xl.rng("A2"), SetType.AUTO, 0)
    assert xl.ws.row_dimensions[2].height >= 15
    xl.Range.set_column_width(xl.rng("A1"), SetType.VALUE, 18)
    assert xl.ws.column_dimensions["A"].width == 18
    xl.Range.set_column_width(xl.rng("A1"), SetType.AUTO, 0)
    assert xl.ws.column_dimensions["A"].width >= 8.43


def test_convert_text_to_number_and_back(xl):
    xl.ws["A1"] = "12"
    xl.ws["A2"] = "3.5"
    xl.ws["A3"] = "1,000"
    xl.ws["A4"] = "abc"
    xl.Range.convert_text_to_number(xl.rng("A1:A4"), xl.rng("Z1"))
    assert xl.ws["A1"].value == 12
    assert xl.ws["A2"].value == 3.5
    assert xl.ws["A3"].value == 1000
    assert xl.ws["A4"].value == "abc"
    xl.Range.convert_number_to_text(xl.rng("A1:A2"))
    assert xl.ws["A1"].value == "12"
    assert xl.ws["A1"].number_format == "@"


def test_add_and_delete_comment(xl):
    xl.Range.add_comment(xl.rng("A1"), "hello")
    assert xl.ws["A1"].comment.text == "hello"
    xl.Range.delete_comment(xl.rng("A1"))
    assert xl.ws["A1"].comment is None
    with pytest.raises(ValueError, match="不存在批注"):
        xl.Range.delete_comment(xl.rng("A1"))


def test_search_and_replace(xl):
    xl.ws["A1"] = "Foo"
    xl.ws["B1"] = "foo bar"
    xl.ws["A2"] = "other"
    hits = xl.Range.search_and_replace(xl.rng("A1:B2"), "foo", case_flag=False)
    assert hits == [{"row": "1", "col": "A"}, {"row": "1", "col": "B"}]
    exact = xl.Range.search_and_replace(xl.rng("A1:B2"), "Foo", exact_match=True, case_flag=True)
    assert exact == [{"row": "1", "col": "A"}]
    first = xl.Range.search_and_replace(xl.rng("A1:B2"), "foo", match_all=False)
    assert first == [{"row": "1", "col": "A"}]
    xl.Range.search_and_replace(xl.rng("A1:B2"), "foo", replace_str="baz", case_flag=False)
    assert "baz" in str(xl.ws["A1"].value).lower()
    assert "baz" in str(xl.ws["B1"].value).lower()


def test_every_method_exists():
    mod = _load_range_mod()
    names = [
        "get_range_data",
        "get_range_color",
        "get_range_size",
        "set_range_data",
        "set_range_type",
        "delete_range",
        "clear_range",
        "copy_range",
        "paste_range",
        "insert_range",
        "merge_range",
        "autofill_range",
        "set_row_height",
        "set_column_width",
        "convert_text_to_number",
        "convert_number_to_text",
        "add_comment",
        "delete_comment",
        "search_and_replace",
    ]
    for name in names:
        assert hasattr(mod.Range, name)
