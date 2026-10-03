import sys
from pathlib import Path

import openpyxl
import pytest
from astronverse.excel import (
    ApplicationType,
    CloseRangeType,
    EditRangeType,
    FileExistenceType,
    ReadRangeType,
    SaveType,
    SaveType_ALL,
    SheetInsertType,
    SheetRangeType,
)
from astronverse.excel.core_openpyxl import OpenpyxlRange, book_of
from astronverse.excel.core_openpyxl.application import Application, get_default_excel_application
from astronverse.excel.core_openpyxl.worksheet import Worksheet
from astronverse.excel.excel import Excel

pytestmark = pytest.mark.skipif(sys.platform == "win32", reason="openpyxl backend is not used on Windows")


@pytest.fixture(autouse=True)
def _quit_app():
    yield
    Application.quit_app(save_changes=False)


def _create(tmp_path: Path, name: str = "demo"):
    return Excel.create_excel(
        file_path=str(tmp_path),
        file_name=name,
        default_application=ApplicationType.EXCEL,
        visible_flag=False,
        exist_handle_type=FileExistenceType.OVERWRITE,
    )


def test_get_default_excel_application():
    assert get_default_excel_application() == ApplicationType.EXCEL


def test_create_edit_read_save_close(tmp_path: Path):
    excel, created_path = _create(tmp_path)
    assert Path(created_path).exists()
    assert excel.get_name() == "demo.xlsx"

    Excel.edit_excel(excel, edit_range=EditRangeType.CELL, start_col="A", start_row="1", value="hello")
    Excel.edit_excel(excel, edit_range=EditRangeType.ROW, start_col="A", start_row="2", value=["a", "b", "c"])

    assert Excel.read_excel(excel, read_range=ReadRangeType.CELL, cell="A1", read_display=False) == "hello"
    row2 = Excel.read_excel(excel, read_range=ReadRangeType.ROW, row="2", read_display=False)
    assert row2[:3] == ["a", "b", "c"]

    Excel.add_excel_worksheet(excel, sheet_name="Data", insert_type=SheetInsertType.LAST)
    names = Excel.get_excel_worksheet_names(excel, sheet_range=SheetRangeType.ALL)
    assert "Data" in names
    Excel.rename_excel_worksheet(excel, source_sheet_name="Data", new_sheet_name="Info")
    names = Worksheet.get_all_worksheet_names(excel)
    assert "Info" in names
    assert "Data" not in names
    Excel.delete_excel_worksheet(excel, del_sheet_name="Info")
    assert "Info" not in Worksheet.get_all_worksheet_names(excel)

    assert Excel.get_excel_row_num(excel) == 2
    assert Excel.get_excel_col_num(excel) == 3
    assert excel.get_first_free_row() == 3

    Excel.save_excel(excel, save_type=SaveType.SAVE)
    Excel.close_excel(close_range_flag=CloseRangeType.ONE, excel=excel, save_type_one=SaveType.ABORT)

    wb = openpyxl.load_workbook(created_path)
    ws = wb.active
    assert ws["A1"].value == "hello"
    assert ws["A2"].value == "a"
    assert ws["C2"].value == "c"
    wb.close()


def test_open_excel_and_get_excel(tmp_path: Path):
    excel, created_path = _create(tmp_path, "opened")
    Excel.edit_excel(excel, edit_range=EditRangeType.CELL, value="x")
    Excel.save_excel(excel, save_type=SaveType.SAVE)
    Excel.close_excel(close_range_flag=CloseRangeType.ONE, excel=excel, save_type_one=SaveType.ABORT)

    excel2 = Excel.open_excel(file_path=created_path, visible_flag=False)
    assert excel2.get_name() == "opened.xlsx"
    assert Excel.read_excel(excel2, read_range=ReadRangeType.CELL, cell="A1", read_display=False) == "x"
    found = Excel.get_excel("opened")
    assert found.get_name() == "opened.xlsx"
    Excel.close_excel(close_range_flag=CloseRangeType.ALL, save_type_all=SaveType_ALL.ABORT)


def test_xls_rejected(tmp_path: Path):
    xls = tmp_path / "old.xls"
    xls.write_bytes(b"not-xlsx")
    with pytest.raises(Exception, match="xls"):
        Excel.open_excel(file_path=str(xls), visible_flag=False)


def test_used_range_trims_empty(tmp_path: Path):
    excel, _ = _create(tmp_path, "used")
    ws = Worksheet.get_worksheet(excel, "", default=1)
    assert Worksheet.get_worksheet_used_range(ws) == (1, 1, 0, 0, "A1")
    Excel.edit_excel(excel, edit_range=EditRangeType.CELL, start_col="B", start_row="2", value="v")
    ws.cell(50, 5)
    start_row, start_col, end_row, end_col, address = Worksheet.get_worksheet_used_range(ws)
    assert (start_row, start_col, end_row, end_col) == (2, 2, 2, 2)
    assert address == "B2"


def test_openpyxl_range_from_address(tmp_path: Path):
    excel, _ = _create(tmp_path, "rng")
    ws = Worksheet.get_worksheet(excel, "", default=1)
    ws["C5"] = 1
    a1 = OpenpyxlRange.from_address(ws, "A1")
    assert a1.address == "A1"
    assert a1.min_row == a1.max_row == a1.min_col == a1.max_col == 1
    area = OpenpyxlRange.from_address(ws, "A1:C3")
    assert area.address == "A1:C3"
    col = OpenpyxlRange.from_address(ws, "A:A")
    assert col.min_col == col.max_col == 1
    row = OpenpyxlRange.from_address(ws, "1:1")
    assert row.min_row == row.max_row == 1
    cells = list(area.cells())
    assert len(cells) == 9
    book = book_of(ws)
    assert book is excel.obj
    assert book_of(book.wb) is book
